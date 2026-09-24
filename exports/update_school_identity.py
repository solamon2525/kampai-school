from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\School คำไผ่\kampai-school")
FILES = [
    ROOT / "exports" / "ชุด-A-แผนวิจัยและการรับรอง-พร้อมใช้งาน.docx",
    ROOT / "public" / "docs" / "classroom-research" / "p4-multiply-race-research-plan-a.docx",
]
FONT = "TH Sarabun New"
LIGHT = "F2F4F7"
NAVY = "17365D"

def font(run, size=16, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rpr.rFonts.set(qn(f"w:{key}"), FONT)
    if color:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string(color)

def set_paragraph(p, text, size=16, bold=False):
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    font(p.add_run(text), size, bold)

def cell_margins(cell):
    pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", 90), ("bottom", 90), ("start", 120), ("end", 120)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    pr.append(mar)

def shade(cell):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT)
    pr.append(shd)

def fill_school_table(table):
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    rows = [
        ("ชื่อสถานศึกษา", "โรงเรียนบ้านคำไผ่"),
        ("กลุ่มเครือข่าย", "กลุ่มเครือข่ายโรงเรียนกุมภวาปี 1"),
        ("ที่ตั้ง", "เลขที่ 159 หมู่ 9 ตำบลเวียงคำ อำเภอกุมภวาปี จังหวัดอุดรธานี 41110"),
        ("สังกัด", "สำนักงานเขตพื้นที่การศึกษาประถมศึกษาอุดรธานี เขต 2"),
        ("ผู้อำนวยการ", "นายสมพิศ แรงน้อย"),
        ("โทรศัพท์", "(โปรดระบุ)"),
        ("อีเมล", "(โปรดระบุ)"),
        ("ระดับที่เปิดสอน", "(โปรดระบุ)"),
        ("แหล่งข้อมูล", "ข้อมูลสถานศึกษาที่ผู้จัดทำยืนยัน เมื่อวันที่ 11 กรกฎาคม 2569"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        for cell in cells:
            cell.text = ""
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cells[0])
        font(cells[0].paragraphs[0].add_run(label), 14, True, NAVY)
        font(cells[1].paragraphs[0].add_run(value), 14)

def update(path):
    doc = Document(path)
    replacements = {
        "สำนักงานเขตพื้นที่การศึกษาประถมศึกษากาฬสินธุ์ เขต 3":
            "กลุ่มเครือข่ายโรงเรียนกุมภวาปี 1 · สำนักงานเขตพื้นที่การศึกษาประถมศึกษาอุดรธานี เขต 2",
        "(นายมกรธวัช แสนสง่า)": "(นายสมพิศ แรงน้อย)",
        "ผู้อำนวยการสถานศึกษา  วิทยฐานะชำนาญการพิเศษ":
            "ผู้อำนวยการโรงเรียนบ้านคำไผ่  (วิทยฐานะ: โปรดระบุ)",
        "ข้อมูลสถานศึกษา (จากหน้าเว็บโรงเรียน)": "ข้อมูลสถานศึกษา (ข้อมูลยืนยันล่าสุด)",
    }
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in replacements:
            set_paragraph(p, replacements[text], 16, text.startswith("("))

    # Document control table.
    for row in doc.tables[0].rows:
        if row.cells[0].text.strip() == "สถานศึกษา":
            set_paragraph(
                row.cells[1].paragraphs[0],
                "โรงเรียนบ้านคำไผ่ · กลุ่มเครือข่ายโรงเรียนกุมภวาปี 1 · สพป.อุดรธานี เขต 2",
                14,
            )

    # Replace the incorrect school profile wholesale instead of retaining stale facts.
    fill_school_table(doc.tables[1])

    doc.core_properties.subject = (
        "ชุด A — แผนวิจัยและการรับรอง โรงเรียนบ้านคำไผ่ "
        "กลุ่มเครือข่ายโรงเรียนกุมภวาปี 1"
    )
    doc.save(path)

for file_path in FILES:
    update(file_path)
    print(file_path)
